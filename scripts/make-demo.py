#!/usr/bin/env python3
"""Builds the App Review demo (server/src/demo/): a made-up story and made-up pictures, files and portraits.
Nothing in it comes from E11's real work: invented clients, invented people, drawn pictures.
Run from the repo root:  python3 scripts/make-demo.py   (needs Pillow, python-docx, python-pptx, openpyxl, ffmpeg)

The story is written with the sample's internal client keys and staff names (so the sample builder can place it);
server/src/demo.js swaps every one of them for the made-up ones below before anything is shown."""
import json, math, os, random, subprocess, tempfile
from PIL import Image, ImageDraw, ImageFilter, ImageFont

OUT = os.path.join(os.path.dirname(__file__), '..', 'server', 'src', 'demo')
os.makedirs(OUT, exist_ok=True)
FONT = '/usr/share/fonts/opentype/inter/Inter-Bold.otf'
FONT_R = '/usr/share/fonts/opentype/inter/Inter-Regular.otf'
FONT_X = '/usr/share/fonts/opentype/inter/InterDisplay-Black.otf' if os.path.exists('/usr/share/fonts/opentype/inter/InterDisplay-Black.otf') else FONT
rnd = random.Random(11)
def font(size, f=FONT): return ImageFont.truetype(f, size)

def grad(w, h, a, b, angle=0.0):
    img = Image.new('RGB', (w, h), a)
    px = img.load()
    ca, cb = [int(a[i:i + 2], 16) for i in (1, 3, 5)], [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    for y in range(0, h, 2):
        for x in range(0, w, 2):
            t = ((x / w) * math.cos(angle) + (y / h) * math.sin(angle) + 1) / 2 if angle else y / h
            t = max(0, min(1, t))
            c = tuple(int(ca[i] + (cb[i] - ca[i]) * t) for i in range(3))
            px[x, y] = c
            if x + 1 < w: px[x + 1, y] = c
            if y + 1 < h: px[x, y + 1] = c; px[min(x + 1, w - 1), y + 1] = c
    return img

def wrap(d, text, f, width):
    words, lines, cur = text.split(), [], ''
    for w in words:
        t = (cur + ' ' + w).strip()
        if d.textlength(t, font=f) > width and cur: lines.append(cur); cur = w
        else: cur = t
    if cur: lines.append(cur)
    return lines

def poster(name, w, h, c1, c2, brand, headline, sub='', motif='circles', fmt='JPEG'):
    img = grad(w, h, c1, c2, 0.9)
    over = Image.new('RGBA', (w, h), (0, 0, 0, 0))
    d = ImageDraw.Draw(over)
    if motif == 'circles':
        for _ in range(7):
            r = rnd.randint(w // 10, w // 3); x = rnd.randint(-r // 2, w); y = rnd.randint(int(h * .1), int(h * .75))
            d.ellipse([x - r, y - r, x + r, y + r], fill=(255, 255, 255, rnd.randint(18, 50)))
    elif motif == 'car':
        y0 = int(h * .52)
        d.rounded_rectangle([w * .12, y0, w * .88, y0 + h * .11], radius=int(h * .05), fill=(255, 255, 255, 230))
        d.rounded_rectangle([w * .28, y0 - h * .07, w * .70, y0 + 10], radius=int(h * .05), fill=(255, 255, 255, 200))
        for cx in (w * .3, w * .7):
            d.ellipse([cx - w * .07, y0 + h * .07, cx + w * .07, y0 + h * .07 + w * .14], fill=(20, 20, 26, 255))
            d.ellipse([cx - w * .03, y0 + h * .07 + w * .04, cx + w * .03, y0 + h * .07 + w * .10], fill=(200, 200, 210, 255))
        d.rectangle([0, y0 + h * .11 + w * .07, w, h], fill=(0, 0, 0, 60))
    elif motif == 'waves':
        for k in range(6):
            pts = [(x, int(h * (.55 + k * .06) + math.sin(x / w * 6.3 + k) * h * .025)) for x in range(0, w + 20, 20)]
            d.line(pts, fill=(255, 255, 255, 70 + k * 20), width=max(3, w // 120))
    elif motif == 'towers':
        for k in range(7):
            bw = w * rnd.uniform(.08, .14); x = w * (.05 + k * .13); top = h * rnd.uniform(.28, .55)
            d.rectangle([x, top, x + bw, h * .82], fill=(255, 255, 255, 60 + k * 18))
            for yy in range(int(top + 20), int(h * .8), 34):
                d.rectangle([x + 10, yy, x + bw - 10, yy + 12], fill=(255, 240, 200, 90))
    elif motif == 'cup':
        cx, cy = w / 2, h * .56
        d.rounded_rectangle([cx - w * .2, cy - h * .12, cx + w * .2, cy + h * .14], radius=40, fill=(255, 255, 255, 235))
        d.ellipse([cx + w * .14, cy - h * .06, cx + w * .32, cy + h * .08], outline=(255, 255, 255, 235), width=max(8, w // 50))
        for k in range(3):
            d.arc([cx - 60 + k * 50, cy - h * .3, cx - 10 + k * 50, cy - h * .16], 200, 340, fill=(255, 255, 255, 160), width=8)
    elif motif == 'bolt':
        d.polygon([(w * .55, h * .25), (w * .32, h * .58), (w * .5, h * .58), (w * .42, h * .82), (w * .7, h * .45), (w * .52, h * .45)], fill=(255, 255, 255, 230))
    elif motif == 'leaf':
        for k in range(9):
            a = k / 9 * math.pi * 2; cx, cy = w / 2 + math.cos(a) * w * .22, h * .55 + math.sin(a) * w * .22
            d.ellipse([cx - w * .09, cy - w * .05, cx + w * .09, cy + w * .05], fill=(255, 255, 255, 110))
        d.ellipse([w / 2 - w * .08, h * .55 - w * .08, w / 2 + w * .08, h * .55 + w * .08], fill=(255, 255, 255, 230))
    img = Image.alpha_composite(img.convert('RGBA'), over.filter(ImageFilter.GaussianBlur(1)))
    d = ImageDraw.Draw(img)
    pad = int(w * .08)
    d.text((pad, pad), brand.upper(), font=font(int(w * .035)), fill=(255, 255, 255, 235))
    hf = font(int(w * .085), FONT_X)
    y = int(h * .13)
    for line in wrap(d, headline, hf, w - pad * 2):
        d.text((pad, y), line, font=hf, fill='white'); y += int(w * .1)
    if sub:
        sf = font(int(w * .034), FONT_R)
        y = int(h * .88)
        for line in wrap(d, sub, sf, w - pad * 2)[:2]:
            d.text((pad, y), line, font=sf, fill=(255, 255, 255, 220)); y += int(w * .045)
    img = img.convert('RGB')
    path = os.path.join(OUT, name)
    img.save(path, fmt, **({'quality': 80, 'optimize': True} if fmt == 'JPEG' else {'optimize': True}))
    return img

# ---------- portraits: drawn, made-up people ----------
SKIN = ['#f1c7a5', '#e0ac85', '#c68c63', '#a86b45', '#8a5634', '#f5d5bd', '#d9a07a']
HAIR = ['#2b1d16', '#4a3020', '#1b1b1f', '#7a4a2a', '#b07a45', '#3a2a22', '#5b5b66']
BGS = [('#7b68ee', '#b3a6ff'), ('#0ea5e9', '#7dd3fc'), ('#14b8a6', '#99f6e4'), ('#f97316', '#fdba74'), ('#e5484d', '#fda4af'), ('#64748b', '#cbd5e1'), ('#a855f7', '#e9d5ff'), ('#0f8a5f', '#86efac')]
def face(name, k, long_hair, beard=False, glasses=False):
    S = 480
    a, b = BGS[k % len(BGS)]
    img = grad(S, S, a, b).convert('RGBA')
    d = ImageDraw.Draw(img)
    skin, hair = SKIN[k % len(SKIN)], HAIR[(k * 3) % len(HAIR)]
    shirt = ['#1f2937', '#334155', '#f8fafc', '#0b3d91', '#7c2d12', '#065f46', '#4c1d95'][k % 7]
    if long_hair: d.rounded_rectangle([S * .27, S * .22, S * .73, S * .78], radius=90, fill=hair)
    d.ellipse([S * .12, S * .70, S * .88, S * 1.25], fill=shirt)
    d.rectangle([S * .43, S * .58, S * .57, S * .76], fill=skin)
    d.ellipse([S * .31, S * .2, S * .69, S * .66], fill=skin)
    d.pieslice([S * .29, S * .14, S * .71, S * .52], 180, 360, fill=hair)
    if not long_hair: d.rectangle([S * .29, S * .3, S * .33, S * .4], fill=hair); d.rectangle([S * .67, S * .3, S * .71, S * .4], fill=hair)
    for ex in (.42, .58):
        d.ellipse([S * ex - 9, S * .41 - 9, S * ex + 9, S * .41 + 9], fill='#1b1b1f')
    if glasses:
        for ex in (.42, .58): d.ellipse([S * ex - 30, S * .41 - 24, S * ex + 30, S * .41 + 24], outline='#111', width=6)
        d.line([S * .42 + 30, S * .41, S * .58 - 30, S * .41], fill='#111', width=6)
    d.arc([S * .43, S * .47, S * .57, S * .57], 20, 160, fill='#7a3b2e', width=6)
    if beard: d.chord([S * .33, S * .38, S * .67, S * .67], 20, 160, fill=hair)
    img = img.filter(ImageFilter.GaussianBlur(.6)).convert('RGB')
    img.save(os.path.join(OUT, name), 'JPEG', quality=82, optimize=True)

PEOPLE = {   # internal id → (made-up name the demo shows, long hair, beard, glasses)
    'u-review': ('Alex Morgan', False, False, True), 'u-nadia': ('Leila Aziz', True, False, False), 'u-lina': ('Lena Farouk', True, False, True),
    'u-lana': ('Maya Rahman', True, False, False), 'u-karim': ('Claire Dubois', True, False, True), 'u-ali': ('Tariq Mansour', False, True, False),
    'u-moon': ('Kai Turner', False, False, False), 'u-maab': ('Nadine Fares', True, False, False), 'u-omar': ('Ravi Menon', False, True, True),
    'u-aish': ('Priya Nair', True, False, False), 'u-new': ('Elena Petrova', True, False, False), 'u-maya': ('Sofia Marin', True, False, False),
    'u-joe': ('Marco Silva', False, True, False), 'u-sara': ('Omar Nasser', False, True, False), 'u-arsal': ('Daniel Reyes', False, False, False),
    'u-smp-free': ('Sam Rivera', False, False, False), 'u-smp-fin-a': ('Rania Haddad', True, False, True), 'u-smp-fin-v': ('Yusuf Karim', False, True, False),
}
for k, (pid, (nm, lh, bd, gl)) in enumerate(PEOPLE.items()): face(f'face-{pid}.jpg', k, lh, bd, gl)

# ---------- the work: made-up campaigns for made-up clients ----------
C = {   # internal client key → (made-up name, colours, motif)
    'Ford & Lincoln': ('Northline Motors', '#0b3d91', '#3b82f6', 'car'), 'KIA': ('Kestrel Auto', '#9f1239', '#f43f5e', 'car'),
    'DEEPAL': ('Voltra EV', '#111827', '#22d3ee', 'bolt'), 'AL TAYER MOTORS': ('Crescent Motors', '#374151', '#9ca3af', 'car'),
    'ALAIN': ('Azure Bay Residences', '#1e3a8a', '#7b68ee', 'towers'), 'HRE': ('Halden Hotels', '#065f46', '#34d399', 'waves'),
    'KOREA 360': ('Seoul Street Market', '#be123c', '#fb7185', 'circles'), 'CPT': ('Coastline Pantry', '#b45309', '#fbbf24', 'cup'),
    'ASV': ('Aster Health', '#0369a1', '#38bdf8', 'leaf'), 'Gebriel Motors': ('Granite Auto Parts', '#0f766e', '#5eead4', 'car'),
    'Bizreh': ('Bloom Café', '#7e22ce', '#f0abfc', 'cup'), 'FORTA': ('Fortis Fitness', '#c2410c', '#fb923c', 'bolt'), 'E11': ('E11', '#334155', '#94a3b8', 'circles'),
}
media_n = [0]
def pic(client, headline, sub='', size=(1080, 1350), fname=None):
    media_n[0] += 1
    key = f'dm{media_n[0]:02d}.jpg'
    name, a, b, motif = C[client]
    poster(key, size[0], size[1], a, b, name, headline, sub, motif)
    short = name.replace(' ', '')
    return {'key': key, 'name': fname or f'{short}_{headline.title().replace(" ", "")[:28]}_{size[0]}x{size[1]}_V0{1 + media_n[0] % 3}.jpg'}

def T(client, title, status, notes, creator, comments, media=(), finance=''):
    return {'key': f'demo{len(TASKS) + 1:02d}', 'client': client, 'title': title, 'status': status, 'finance': finance, 'notes': notes, 'creator': creator,
            'comments': [{'who': w, 'text': t, 'media': list(m), 'reacts': r, 'replies': 0} for (w, t, m, r) in comments], 'media': list(media)}
TASKS = []
def add(*a, **k): TASKS.append(T(*a, **k))

N, K, V, CR, AZ, HD, SS, CP, AH, GA, BC, FF = 'Ford & Lincoln', 'KIA', 'DEEPAL', 'AL TAYER MOTORS', 'ALAIN', 'HRE', 'KOREA 360', 'CPT', 'ASV', 'Gebriel Motors', 'Bizreh', 'FORTA'
add(N, 'Northline | Winter Drive Campaign – Key Visual', 'WIP', 'Hero key visual for the Winter Drive campaign. Mountain road at dusk, the new SUV in front, warm headlights.\nHeadline: Every road, every season.\nSizes: 1080x1350, 1080x1920, 1920x1080.', 'u-nadia',
    [('u-nadia', 'Hi @Joe Hajj, brief and references are in the description. First direction by Thursday please.', [], 0), ('u-joe', 'On it. Sharing two looks tomorrow morning.', [], 1)],
    [pic(N, 'Every road, every season', 'Winter Drive · test drive this weekend')], 'PO Received')
add(N, 'Northline | Service Offer – Social Pack', 'INTERNAL REVIEW', 'Four posts for the winter service offer: 20% off the 10-point check, free car wash.', 'u-nadia',
    [('u-sara', 'Service pack ready for review, all four sizes attached.', [], 0)], [pic(N, 'Winter service check', '20% off this month'), pic(N, 'Free car wash', 'With every service booking')], 'LPO Pending')
add(N, 'Northline | Showroom Opening – Invitation EDM', 'AMENDMENTS', 'English + Arabic invitation for the new showroom opening. RSVP button links to the event page.', 'u-nadia',
    [('u-karim', 'The date line is too small, make it the second thing you read.', [], 0), ('u-sara', 'Noted, updating now.', [], 0)], [pic(N, 'You are invited', 'Showroom opening · Saturday 7 PM')], 'PO Received')
add(N, 'Northline | Weekend Test Drive – Reel', 'EDIT VIDEO', '15s reel from the test drive shoot. Fast cuts, end card with the booking link.', 'u-nadia',
    [('u-moon', 'Rough cut is on the NAS, music still temporary.', [], 0)], [], 'Quote to be sent')
add(N, 'Northline | Owners Club – Newsletter', 'COMPLETED', 'Monthly owners club newsletter with the road trip story.', 'u-nadia', [('u-joe', 'Final files attached.', [], 1)], [pic(N, 'Owners Club road trip', 'November edition')], 'PO Received')
add(K, 'Kestrel | Family Cars Campaign – Static', 'WIP', 'Static posts for the family cars campaign: seven seats, room for everything.', 'u-lina',
    [('u-lina', '@Yousef Al Hwij please keep the family in the shot, not only the car.', [], 0), ('u-sara', 'Got it, reworking the layout.', [], 0)], [pic(K, 'Room for everyone', 'Seven seats. Zero compromises.')], 'PO Received')
add(K, 'Kestrel | Year End Sale – Banners', 'CTR', 'Website and showroom banners for the year end sale.', 'u-lina', [('u-ali', 'All banner sizes approved by the CD, attached for the client.', [], 1)],
    [pic(K, 'Year end sale', 'Offers on every model', (1920, 1080))], 'To be invoiced')
add(K, 'Kestrel | Hybrid Launch – Storyboard', 'STORYBOARD', 'Storyboard for the 30s hybrid launch film. City at night, quiet electric start.', 'u-lina', [('u-ali', 'Frames 1–6 drafted, sharing after lunch.', [], 0)], [], 'Quote to be sent')
add(K, 'Kestrel | Safety Week – Carousel', 'AMENDMENTS', 'Five-slide carousel on the five-star safety rating.', 'u-lina',
    [('u-karim', 'Slide 3 icon style doesn’t match the others.', [], 0)], [pic(K, 'Five stars for safety', 'Swipe to see why')], 'LPO Pending')
add(V, 'Voltra | City Charge – Launch Video', 'SHOOT', 'Shoot day at the marina: the car charging, the driver walking to a café. 2 locations, golden hour.', 'u-lina',
    [('u-moon', 'Location permits confirmed for Saturday.', [], 2)], [], 'PO Received')
add(V, 'Voltra | Range Myths – Social Series', 'WIP', 'Four short posts answering common range questions.', 'u-lina', [('u-joe', 'Copy for posts 1 and 2 is in the doc.', [], 0)],
    [pic(V, 'Range myths, busted', 'Part 1 · How far can it go?')], 'Retainer')
add(V, 'Voltra | Pre-Book – Landing Banner', 'CTR', 'Hero banner for the pre-book landing page.', 'u-lina', [('u-maab', 'Banner approved, ready for the client.', [], 0)], [pic(V, 'Pre-book yours today', 'Limited first edition', (1920, 1080))], 'To be invoiced')
add(CR, 'Crescent | Group Brand Film – Script', 'BRAINSTORM', 'Brainstorm for the group brand film: 40 years of moving people.', 'u-lina', [], [], 'Quote to be sent')
add(AZ, 'Azure Bay | Tower Launch – Brochure', 'WIP', 'Twelve-page brochure for the new tower: floor plans, amenities, location map.', 'u-lana',
    [('u-lana', '@Ali Raza the floor plans came in from the client, they are in the Files folder.', [], 0), ('u-ali', 'Thanks, adding them today.', [], 0)], [pic(AZ, 'Living above the bay', 'Tower launch brochure')], 'PO Received')
add(AZ, 'Azure Bay | Open House – Reel', 'INTERNAL REVIEW', 'Open house invitation reel, 9:16.', 'u-lana', [('u-maya', 'Reel ready for review.', [], 0)], [pic(AZ, 'Open house this Friday', '', (1080, 1920))], 'Retainer')
add(AZ, 'Azure Bay | Amenities – Carousel', 'COMPLETED', 'Carousel showing the pool, gym and kids club.', 'u-lana', [('u-ali', 'Final carousel attached.', [], 1)], [pic(AZ, 'Everything you need, downstairs')], 'PO Received')
add(HD, 'Halden | Staycation Offer – Stories', 'WIP', 'Story set for the weekend staycation offer.', 'u-lana', [('u-maab', 'First two frames done.', [], 0)], [pic(HD, 'Your weekend escape', 'Stay 2 nights, get the 3rd free', (1080, 1920))], 'Retainer')
add(HD, 'Halden | Festive Brunch – Poster', 'AMENDMENTS', 'Poster for the festive brunch, A2 and social.', 'u-lana', [('u-karim', 'Make the price clearer and drop the second font.', [], 0)], [pic(HD, 'Festive brunch', 'Every Friday in December')], 'LPO Pending')
add(SS, 'Seoul Street | Food Week – Reel', 'WIP', 'Tasting reel at the food stalls.', 'u-nadia', [('u-aish', 'Shot list shared with the team.', [], 0)], [pic(SS, 'Food week is here', '', (1080, 1920))], 'Retainer')
add(SS, 'Seoul Street | Monthly Content Plan', 'WIP', 'Content plan for next month: events, giveaways, weekend activities.', 'u-aish', [], [], 'Retainer')
add(CP, 'Coastline | New Menu – Static Posts', 'CTR', 'Three posts for the new breakfast menu.', 'u-aish', [('u-aish', 'Approved by the CD, sharing with the client now.', [], 0)], [pic(CP, 'Good mornings start here', 'New breakfast menu')], 'Retainer')
add(CP, 'Coastline | Weekend Market – Story', 'NEW', 'Story announcing the weekend farmers market at the bakery.', 'u-aish', [], [], 'Retainer')
add(CP, 'Coastline | Monthly Content Plan', 'WIP', 'Next month’s plan: menu highlights, behind the scenes, customer stories.', 'u-maya', [], [], 'Retainer')
add(AH, 'Aster Health | Wellness Month – Carousel', 'WIP', 'Carousel with five everyday wellness tips.', 'u-maya', [('u-omar', 'Illustrations in progress.', [], 0)], [pic(AH, 'Five small habits', 'For a healthier you')], 'Retainer')
add(AH, 'Aster Health | Clinic Opening – Video', 'ON HOLD', 'Video for the new clinic opening, waiting for the client’s date.', 'u-maya', [('u-maya', 'On hold until the client confirms the opening date.', [], 0)], [], 'Retainer')
add(AH, 'Aster Health | Flu Season – Posts', 'COMPLETED', 'Reminder posts for flu vaccinations.', 'u-maya', [('u-omar', 'Final posts attached.', [], 0)], [pic(AH, 'Get your flu shot', 'Walk in, no appointment')], 'Retainer')
add(GA, 'Granite | Trade Show – Mini Stand Static', 'WIP', 'We’re not exhibiting, so we built our trade show stand in miniature from our own parts.\nHeadline: No stand this year. So we built our own.', 'u-aish',
    [('u-joe', 'Hello @Aishwarya Sanjay, mini stand static attached.', [], 0)], [pic(GA, 'No stand this year. So we built our own.')], 'Retainer')
add(GA, 'Granite | Brake Pads – Product Posts', 'INTERNAL REVIEW', 'Product posts for the new ceramic brake pads.', 'u-aish', [('u-joe', 'Submitted for review.', [], 0)], [pic(GA, 'Stop with confidence', 'New ceramic brake pads')], 'Retainer')
add(GA, 'Granite | Workshop Tips – Reel', 'NEW', 'Short reel: three quick checks before a long drive.', 'u-aish', [], [], 'Retainer')
add(BC, 'Bloom Café | Autumn Drinks – Posts', 'AMENDMENTS', 'Posts for the autumn drinks menu.', 'u-maya', [('u-karim', 'Warmer colours please, it reads like summer.', [], 0)], [pic(BC, 'Autumn in a cup', 'Spiced latte is back')], 'Retainer')
add(FF, 'Fortis | New Year Membership – Campaign', 'WIP', 'Campaign for the new year membership offer: posts, stories and an EDM.', 'u-nadia', [('u-ali', 'Moodboard shared in the folder.', [], 0)], [pic(FF, 'New year, stronger you', 'Join this month, no joining fee')], 'Quote to be sent')
add('E11', 'E11 | Agency Showreel 2026', 'WIP', 'Our own showreel with the best work of the year.', 'u-lana', [('u-moon', 'Collecting the selects this week.', [], 0)], [], '')
add(K, 'Kestrel | Dealer Training – Slides', 'COMPLETED', 'Slides for the dealer training day.', 'u-lina', [('u-sara', 'Final deck attached.', [], 0)], [pic(K, 'Dealer training day', '', (1920, 1080))], 'PO Received')

json.dump({'tasks': TASKS}, open(os.path.join(OUT, 'demo-story.json'), 'w'), ensure_ascii=False, indent=1)

# ---------- the sample's general files, made up (same names as server/src/sample-files) ----------
poster('ford-key-visual.jpg', 1080, 1350, '#0b3d91', '#3b82f6', 'Northline Motors', 'Built for the long way home', 'Key visual · V03', 'car')
poster('alain-tower-render.jpg', 1080, 1350, '#1e3a8a', '#7b68ee', 'Azure Bay Residences', 'Tower render', 'Evening view from the marina', 'towers')
poster('korea360-recap-cover.jpg', 1080, 1350, '#be123c', '#fb7185', 'Seoul Street Market', 'Event recap', 'Food week highlights', 'circles')
poster('story-frame-1.png', 540, 960, '#065f46', '#34d399', 'Halden Hotels', 'Frame 1', 'Pool at sunrise', 'waves', fmt='PNG')

def pdf(name, title, lines, color):
    pages = []
    for p in range(2):
        img = Image.new('RGB', (1240, 1754), 'white'); d = ImageDraw.Draw(img)
        d.rectangle([0, 0, 1240, 160], fill=color)
        d.text((80, 52), title if p == 0 else title + ' (continued)', font=font(52), fill='white')
        y = 240
        for ln in lines[p * 9:(p + 1) * 9]:
            d.text((80, y), ln, font=font(30, FONT_R), fill='#222'); y += 64
        d.text((80, 1660), 'Made-up document for the App Review demo', font=font(22, FONT_R), fill='#999')
        pages.append(img)
    pages[0].save(os.path.join(OUT, name), save_all=True, append_images=pages[1:], resolution=150)
pdf('creative-brief.pdf', 'Creative brief', ['Client: Northline Motors', 'Campaign: Winter Drive', 'Objective: test drive bookings', 'Audience: families, 30–45', 'Key message: every road, every season',
    'Deliverables: key visual, 4 posts, reel', 'Tone: warm, confident', 'Timing: 3 weeks', 'Budget: approved', 'Mandatories: logo end card', 'Arabic and English copy', 'Booking link on every post'], '#7b68ee')
pdf('invoice-expense.pdf', 'Invoice  #INV-2041', ['Bright Lights Studio Rental LLC', 'Studio rental, 2 days', 'Lighting kit and grip', 'Subtotal: AED 4,000.00', 'VAT 5%: AED 200.00', 'Total: AED 4,200.00', 'Paid by bank transfer'], '#0ea5e9')
pdf('invoice-subscription.pdf', 'Subscription invoice', ['Design Cloud Pro, team plan', '12 seats, monthly', 'Total: AED 1,250.00', 'Next billing in one month'], '#14b8a6')
pdf('receipt-reimbursement.pdf', 'Receipt', ['Harbour Fuel Station', 'Petrol 98, 45.2 L', 'Total: AED 180.00', 'Card ending 0000'], '#f97316')

from docx import Document
doc = Document(); doc.add_heading('Copy deck · Winter Drive', 0)
for h, t in [('Headline', 'Every road, every season.'), ('Post 1', 'Snow or sand, it’s ready. Book your test drive this weekend.'), ('Post 2', 'Warm seats, long drives, happy kids.'), ('End card', 'Northline Motors · Book a test drive')]:
    doc.add_heading(h, 2); doc.add_paragraph(t)
doc.save(os.path.join(OUT, 'copy-deck.docx'))
from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation()
for title, body in [('Winter Drive', 'Campaign presentation · Northline Motors'), ('The idea', 'Every road, every season'), ('Deliverables', 'Key visual · 4 posts · 15s reel · EDM'), ('Timeline', 'Week 1 concepts · Week 2 production · Week 3 launch')]:
    s = prs.slides.add_slide(prs.slide_layouts[1]); s.shapes.title.text = title; s.placeholders[1].text = body
prs.save(os.path.join(OUT, 'client-presentation.pptx'))
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = 'Media plan'
ws.append(['Week', 'Channel', 'Format', 'Budget (AED)'])
for r in [[1, 'Instagram', 'Reel', 3000], [1, 'TikTok', 'Reel', 2000], [2, 'Instagram', 'Carousel', 1500], [2, 'LinkedIn', 'Post', 800], [3, 'YouTube', 'Pre-roll', 4000]]: ws.append(r)
wb.save(os.path.join(OUT, 'media-plan.xlsx'))

# a short made-up video: three drawn frames
with tempfile.TemporaryDirectory() as tmp:
    for i, (h, motif) in enumerate([('Coming soon', 'bolt'), ('Charge in 20 minutes', 'waves'), ('Pre-book today', 'circles')]):
        img = poster('_f.jpg', 720, 1280, '#111827', '#22d3ee', 'Voltra EV', h, '', motif)
        img.save(os.path.join(tmp, f'f{i}.jpg'))
    os.remove(os.path.join(OUT, '_f.jpg'))
    subprocess.run(['ffmpeg', '-y', '-loglevel', 'error', '-framerate', '1/2', '-i', os.path.join(tmp, 'f%d.jpg'), '-c:v', 'libx264', '-r', '25', '-pix_fmt', 'yuv420p', '-crf', '30',
                    '-movflags', '+faststart', os.path.join(OUT, 'teaser-rough-cut.mp4')], check=True)
print('demo:', len(TASKS), 'tasks,', media_n[0], 'pictures')
