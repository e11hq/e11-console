"""Content plans (trial, test copy): the list of clients for a month, starting a plan, the month calendar with occasions,
adding and editing posts with references, skipping an occasion, who can edit, the Social Manager and the client's Social
people, and the phone."""
import asyncio, os, subprocess, tempfile, time, urllib.request, urllib.parse, json
from playwright.async_api import async_playwright
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
PORT = 18331; BASE = f'http://127.0.0.1:{PORT}'; ADMIN = 'haris@e11-digital.com'
EXEC, MGR, AM = 'aishwarya@e11-digital.com', 'amelia@e11-digital.com', 'areej@e11-digital.com'
OUT = os.environ.get('SHOTS') or tempfile.mkdtemp(); fails = []
def check(c, m):
    print(('PASS ' if c else 'FAIL ') + m)
    if not c: fails.append(m)
def api(who, p, body=None):
    req = urllib.request.Request(BASE + p, method='POST' if body is not None else 'GET', data=json.dumps(body).encode() if body is not None else None, headers={'X-Dev-Email': who, 'Content-Type': 'application/json'})
    return json.loads(urllib.request.urlopen(req).read())
async def main():
    env = {**os.environ, 'DATA_DIR': tempfile.mkdtemp(), 'AUTH_DEV': '1', 'ADMIN_EMAIL': ADMIN, 'PORT': str(PORT), 'WEB_DIR': os.path.join(ROOT, 'web'), 'APP_VERSION': '1.6.5-test.1'}
    srv = subprocess.Popen(['node', os.path.join(ROOT, 'server/src/index.js')], env=env, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        for _ in range(50):
            try: urllib.request.urlopen(BASE + '/api/health'); break
            except Exception: time.sleep(0.1)
        api(ADMIN, '/api/me'); api(ADMIN, '/api/sample/reset', {})
        async with async_playwright() as pw:
            b = await pw.chromium.launch(executable_path='/opt/pw-browsers/chromium-1194/chrome-linux/chrome')
            async def login(email, vp=(1440, 900)):
                ctx = await b.new_context(viewport={'width': vp[0], 'height': vp[1]}, base_url=BASE, has_touch=vp[0] < 500)
                await ctx.add_cookies([{'name': 'e11_dev_email', 'value': email, 'url': BASE}])
                pg = await ctx.new_page(); errs = []; pg.on('pageerror', lambda e: errs.append(str(e)))
                await pg.goto('/'); await pg.wait_for_timeout(3500); await pg.mouse.click(5, 5); await pg.wait_for_timeout(600)
                return ctx, pg, errs
            # ---------- the Social Executive ----------
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            check(await pg.locator('#nav-plans').is_visible(), 'Content plans in the sidebar (test site)')
            await pg.locator('#nav-plans').click(); await pg.wait_for_timeout(500)
            rows = await pg.locator('.pl-row:not(.tt-head)').count()
            check(rows == 2, f'Aishwarya sees the 2 clients she is Social for ({rows})')
            check('Trial' in await pg.locator('#title-area').inner_text(), 'marked as a trial')
            kr = pg.locator('.pl-row', has_text='KOREA 360')
            check('Open' in await kr.inner_text() and await kr.locator('.pl-st.review').count() == 1, 'KOREA 360 has a skeleton, waiting for approval')
            await pg.screenshot(path=os.path.join(OUT, 'plans-list.png'))
            # start one for ASV
            await pg.locator('.pl-row', has_text='ASV').locator('[data-pl="start"]').click(); await pg.wait_for_timeout(600)
            check(await E("ui.view") == 'plan' and await pg.locator('.pl-who b').inner_text() == 'ASV', 'Start skeleton opens ASV’s new skeleton')
            check(await pg.locator('.pl-cell:not(.out)').count() >= 28, 'the month calendar')
            d = await E("Plans.planById(ui.planId).month + '-15'")
            await pg.locator(f'.pl-cell[data-day="{d}"]').hover()
            await pg.locator(f'.pl-cell[data-day="{d}"] [data-pl="new-post"]').click(); await pg.wait_for_timeout(300)
            check(await pg.locator('.pl-dlg').count() == 1, 'the + on a day opens a new post')
            await pg.fill('#pp-title', 'Sportage launch teaser')
            await pg.locator('[data-fmt="reel"]').click()
            await pg.locator('[data-ch="tt"]').click()
            await pg.fill('#pp-desc', 'Teaser for the new Sportage.')
            await pg.fill('#pp-en', 'Something new is coming.')
            await pg.fill('#pp-ar', 'شيء جديد قادم.')
            await pg.locator('[data-pp="link"]').click(); await pg.fill('#pp-link', 'www.instagram.com/kia'); await pg.keyboard.press('Enter'); await pg.wait_for_timeout(200)
            check(await pg.locator('.pl-ref-link').count() == 1, 'a reference link')
            async with pg.expect_file_chooser() as fc:
                await pg.locator('[data-pp="upload"]').click()
            await (await fc.value).set_files(os.path.join(ROOT, 'server/src/sample-files/clickup/cu09.jpg')); await pg.wait_for_timeout(2500)
            check(await pg.locator('.pl-ref-img img').count() == 1, 'an uploaded picture shows')
            bb = await pg.locator('.pl-dlg').bounding_box()
            check(bb and bb['y'] >= 0 and bb['y'] + bb['height'] <= 900, f'the pop-up fits the window {bb}')
            await pg.screenshot(path=os.path.join(OUT, 'plans-post.png'))
            await pg.locator('.pl-dlg [type="submit"]').click(); await pg.wait_for_timeout(800)
            chip = pg.locator(f'.pl-cell[data-day="{d}"] .pl-chip')
            check(await chip.count() == 1 and 'Reel' in await chip.inner_text(), 'the post is on its day, as a Reel')
            saved = await asyncio.to_thread(lambda: [p for p in api(ADMIN, '/api/snapshot')['data']['plans'] if p['projectId'] == 'p-asv'])
            check(len(saved) == 1 and saved[0]['posts'][0]['title'] == 'Sportage launch teaser' and saved[0]['posts'][0]['channels'] == ['ig', 'tt'] and len(saved[0]['posts'][0]['refs']) == 2, 'saved on the server with its channels and references')
            await chip.click(); await pg.wait_for_timeout(300)
            await pg.fill('#pp-title', 'Sportage teaser'); await pg.locator('.pl-dlg [type="submit"]').click(); await pg.wait_for_timeout(500)
            check('Sportage teaser' in await pg.locator(f'.pl-cell[data-day="{d}"]').inner_text(), 'editing a post')
            await pg.locator('[data-pl="tab"][data-v="list"]').click(); await pg.wait_for_timeout(300)
            check(await pg.locator('.pl-post').count() == 1 and await pg.locator('.pl-post .pl-thumb img').count() == 1, 'the list view with the picture')
            await pg.locator('[data-pl="tab"][data-v="calendar"]').click(); await pg.wait_for_timeout(200)
            # occasions this month (on ASV's new skeleton) and skipping one
            miss = pg.locator('.pl-occ.miss').first
            if await miss.count():
                name = await miss.locator('b').inner_text()
                await miss.locator('[data-pl="skip"]').click(); await pg.wait_for_timeout(500)
                check(await pg.locator('.pl-occ.skip', has_text=name).count() == 1, f'skipping “{name}” for ASV')
                check(await E("(proj('p-asv').planSkips || []).length") == 1, 'remembered on the client (next year too)')
                await pg.locator('.pl-occ.skip', has_text=name).locator('[data-pl="unskip"]').click(); await pg.wait_for_timeout(400)
                check(await pg.locator('.pl-occ.miss', has_text=name).count() == 1, 'undo skip')
                await pg.locator('.pl-occ.miss', has_text=name).locator('[data-pl="occ-post"]').click(); await pg.wait_for_timeout(300)
                check(await pg.input_value('#pp-title') == name and await pg.input_value('#pp-occ') != '', 'Plan a post: the occasion is filled in')
                await pg.locator('.pl-dlg-top [data-pp="close"]').click(); await pg.wait_for_timeout(200)
            # send ASV's skeleton: the occasion warning, then to the Social Manager
            await pg.locator('[data-pl="send"]').click(); await pg.wait_for_timeout(300)
            dl = pg.locator('.pl-send-dlg')
            check(await dl.locator('[data-to]').count() == 3, 'send to: Social Manager, Carine or Both')
            if await pg.locator('.pl-occ.miss').count():
                check(await dl.locator('.pl-warn').count() == 1 and 'Send anyway' in await dl.locator('[data-r="1"]').inner_text(), 'warns about occasions not planned: Send anyway')
            await pg.screenshot(path=os.path.join(OUT, 'plans-send.png'))
            await dl.locator('[data-r="1"]').click(); await pg.wait_for_timeout(600)
            check(await pg.locator('.pl-st.review').count() == 1 and await pg.locator('.pl-banner.review').count() == 1, 'waiting for approval')
            check(await pg.locator('[data-pl="new-post"]').count() == 0, 'locked while it waits')
            # KOREA 360 (sample): already waiting for the Social Manager
            await pg.locator('#nav-plans').click(); await pg.wait_for_timeout(300)
            await pg.locator('.pl-row', has_text='KOREA 360').click(); await pg.wait_for_timeout(500)
            n = await pg.locator('.pl-chip').count()
            check(n >= 6, f'KOREA 360’s sample posts on the calendar ({n})')
            check(await pg.locator('.pl-banner.review').count() == 1, 'KOREA 360 waits for approval')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # ---------- Stage 2: the Social Manager looks at KOREA 360 ----------
            ctx, pg, errs = await login(MGR)
            E = pg.evaluate
            check(await pg.locator('#plans-count').inner_text() == '2', f'the sidebar shows 2 skeletons waiting ({await pg.locator("#plans-count").inner_text()})')
            await E("go('inbox')"); await pg.wait_for_timeout(400)
            row = pg.locator('.ib-row', has_text='KOREA 360')
            full = (await row.inner_text()) + ' ' + ((await row.locator('.ib-mi').get_attribute('title')) or '') if await row.count() == 1 else ''
            check('for your approval' in full, 'Inbox: the skeleton sent for approval (all of it on hover when the line is cut)')
            await row.click(); await pg.wait_for_timeout(600)
            check(await E("ui.view") == 'plan' and await pg.locator('.pl-who b').inner_text() == 'KOREA 360', 'the update opens the skeleton')
            check(await pg.locator('[data-pl="approve"]').count() == 1 and await pg.locator('[data-pl="send-back"]').count() == 1, 'Approve skeleton / Send back for changes')
            await pg.locator('.pl-chip').first.click(); await pg.wait_for_timeout(300)
            await pg.locator('[data-ck="changes"]').click(); await pg.fill('#pp-cknote', 'Use the new key visual')
            await pg.screenshot(path=os.path.join(OUT, 'plans-check.png'))
            await pg.locator('.pl-dlg [type="submit"]').click(); await pg.wait_for_timeout(400)
            check(await pg.locator('.pl-chip.chg').count() == 1, 'the post is marked for changes')
            await pg.locator('[data-pl="approve"]').click(); await pg.wait_for_timeout(300)
            check(await E("Plans.planById('pl-smp-p-korea').status") == 'review', 'can’t approve with a post marked for changes')
            await pg.locator('[data-pl="send-back"]').click(); await pg.wait_for_timeout(200)
            await pg.fill('#sb-note', 'Nearly there'); await pg.locator('.pl-send-dlg [data-r="1"]').click(); await pg.wait_for_timeout(600)
            check(await pg.locator('.pl-st.changes').count() == 1, 'sent back for changes')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # Aishwarya makes the change and sends it to Carine
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            await E("go('inbox')"); await pg.wait_for_timeout(400)
            row = pg.locator('.ib-row', has_text='KOREA 360')
            check(await row.count() == 1 and 'asked for changes' in await row.inner_text(), 'Inbox: changes asked')
            await row.click(); await pg.wait_for_timeout(600)
            check(await pg.locator('.pl-banner.changes').count() == 1 and 'Nearly there' in await pg.locator('.pl-banner').inner_text(), 'the note on the skeleton')
            await pg.locator('.pl-chip.chg').click(); await pg.wait_for_timeout(300)
            check('Use the new key visual' in await pg.locator('.pl-check.chg').inner_text(), 'the note on the post')
            await pg.fill('#pp-desc', 'With the new key visual.'); await pg.locator('.pl-dlg [type="submit"]').click(); await pg.wait_for_timeout(400)
            await pg.locator('[data-pl="send"]').click(); await pg.wait_for_timeout(300)
            await pg.locator('.pl-send-dlg [data-to="cd"]').click(); await pg.locator('.pl-send-dlg [data-r="1"]').click(); await pg.wait_for_timeout(600)
            check(await E("Plans.planById('pl-smp-p-korea').review.to") == 'cd' and await E("Plans.planById('pl-smp-p-korea').round") == 2, 'sent again to Carine (Round 2)')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # Carine approves
            ctx, pg, errs = await login('carine@e11-digital.com')
            E = pg.evaluate
            await E("go('inbox')"); await pg.wait_for_timeout(400)
            await pg.locator('.ib-row', has_text='KOREA 360').click(); await pg.wait_for_timeout(600)
            await pg.locator('[data-pl="approve"]').click(); await pg.wait_for_timeout(300)
            await pg.locator('.dialog [data-r="1"]').click(); await pg.wait_for_timeout(600)
            check(await pg.locator('.pl-st.approved').count() == 1 and await pg.locator('.pl-banner.approved').count() == 1, 'Carine approves the skeleton')
            await pg.screenshot(path=os.path.join(OUT, 'plans-approved.png'))
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            note = await asyncio.to_thread(lambda: [n for n in api(EXEC, '/api/snapshot')['data']['notifications'] if n.get('planId') == 'pl-smp-p-korea' and n['kind'] == 'plan-ok'])
            check(len(note) == 1, 'Aishwarya is told it’s approved')
            # ---------- Stage 3: Aishwarya creates the tasks ----------
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            await E("Plans.open('pl-smp-p-korea')"); await pg.wait_for_timeout(500)
            check(await pg.locator('[data-pl="make-tasks"]').count() == 1, 'Create tasks on the approved skeleton')
            await pg.locator('[data-pl="make-tasks"]').click(); await pg.wait_for_timeout(300)
            n = await pg.locator('.mt-row').count()
            check(n >= 6, f'one row per post ({n})')
            await pg.select_option('#mt-fmt', 'reel')
            await pg.locator('[data-mt="bulk-who"]').click(); await pg.wait_for_timeout(200)
            await pg.locator('.mt-menu .mi', has_text='Joe Hajj').click(); await pg.wait_for_timeout(200)
            await pg.mouse.click(5, 5) if False else None
            await pg.keyboard.press('Escape'); await E("closeMenus()")
            reels = await E("Plans.planById('pl-smp-p-korea').posts.filter(x => x.format === 'reel').length")
            check(await pg.locator('.mt-row .mt-who', has_text='Joe').count() == reels, f'all reels to Joe in one go ({reels})')
            await pg.locator('.mt-row [data-mt="ck"]').last.click()
            await pg.screenshot(path=os.path.join(OUT, 'plans-make-tasks.png'))
            await pg.locator('.mt-dlg [data-r="1"]').click(); await pg.wait_for_timeout(1500)
            made = await E("db.tasks.filter(t => t.planId === 'pl-smp-p-korea')")
            check(len(made) == n - 1, f'{n - 1} tasks made (one left out)')
            t0 = next((t for t in made if t['attachments']), None)
            check(all(t['projectId'] == 'p-korea' and t['golive'] and 'From the content plan' in t['notesHtml'] for t in made), 'each task has the client, go-live and the skeleton’s details')
            check(t0 is not None and t0['attachments'][0]['path'].startswith('attachments/' + t0['id'] + '/'), 'reference files go into the task')
            check(await pg.locator('[data-pl="make-tasks"]').count() == 1, 'Create tasks stays for the post left out')
            check('In production' in await pg.locator('.pl-chip').first.inner_text(), 'the calendar shows each post’s progress')
            tid = made[0]['id']
            check(await E(f"(() => {{ const t = task('{tid}'); return [can('editDates', t), can('assign', t), can('deleteTask', t), can('createTask')]; }})()") == [True, True, True, False], 'Social can do what an AM can on these tasks (not create others)')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            am_note = await asyncio.to_thread(lambda: [n for n in api(AM, '/api/snapshot')['data']['notifications'] if n.get('planId') == 'pl-smp-p-korea' and n['kind'] == 'plan-tasks'])
            check(len(am_note) == 1, 'the client’s AM is told')
            joe = await asyncio.to_thread(lambda: [n for n in api('joe@e11-digital.com', '/api/snapshot')['data']['notifications'] if n['kind'] == 'assigned' and n.get('taskId') in [t['id'] for t in made]])
            jn = len([t for t in made if 'u-joe' in t['assignees']])
            check(jn >= 1 and len(joe) == jn, f'Joe is told about his {jn} tasks ({len(joe)})')
            # the CD approves a task's files: the finished work comes into the post by itself
            ctx, pg, errs = await login('carine@e11-digital.com')
            E = pg.evaluate
            tid = t0['id']
            await E(f"""(() => {{ const t = task('{tid}'); const a = t.attachments[0]; t.reviews = (t.reviews || []).concat({{ id: uid(), by: t.assignees[0] || 'u-joe', files: [a.id], at: nowISO(), state: 'approved', decidedAt: nowISO(), decidedBy: currentUserId, round: 1, epoch: t.reviewEpoch || 0 }}); syncCD(t); if (!t.cdStatus) t.cdStatus = 'approved'; persist(0); }})()"""); await pg.wait_for_timeout(800)
            await E("Plans.open('pl-smp-p-korea')"); await pg.wait_for_timeout(500)
            pid = t0['planPostId']
            chip = pg.locator(f'.pl-chip[data-id="{pid}"]')
            check('Finished' in await chip.inner_text(), 'Finished on the calendar once the CD approves')
            await chip.click(); await pg.wait_for_timeout(300)
            check(await pg.locator('.pl-task-box .pl-ref-img').count() == 1, 'the finished work shows in the post')
            await pg.screenshot(path=os.path.join(OUT, 'plans-finished.png'))
            check('1 of' in await pg.locator('.pl-banner').inner_text(), 'the skeleton counts finished posts')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # ---------- Stage 4: the client round ----------
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            await E("Plans.open('pl-smp-p-korea')"); await pg.wait_for_timeout(500)
            await pg.locator('[data-pl="make-tasks"]').click(); await pg.wait_for_timeout(300)
            await pg.locator('.mt-dlg [data-r="1"]').click(); await pg.wait_for_timeout(1200)
            ids = await E("Plans.planById('pl-smp-p-korea').posts.map(x => x.taskId)")
            check(all(ids), 'every post has its task now')
            await ctx.close()
            reel = await asyncio.to_thread(lambda: next(t['id'] for t in api(ADMIN, '/api/snapshot')['data']['tasks'] if t.get('planId') == 'pl-smp-p-korea' and t['title'] == 'K-pop night'))
            import base64
            vid = await asyncio.to_thread(lambda: api('carine@e11-digital.com', '/api/host/writeChunk', {'taskId': reel, 'name': 'kpop-night.mp4', 'b64': base64.b64encode(open(os.path.join(ROOT, 'server/src/sample-files/clickup/cu52.mp4'), 'rb').read()).decode()}))
            # the CD approves every task's work (a picture each, the video on the reel)
            APPROVE = """(async (vid) => { const ts = db.tasks.filter(t => t.planId === 'pl-smp-p-korea' && !t.done); const src = db.tasks.find(t => t.planId === 'pl-smp-p-korea' && t.attachments.length).attachments[0];
              for (const t of ts) { let a;
                if (vid && t.title === 'K-pop night') a = { id: uid(), name: vid.name, path: vid.path, size: vid.size, addedAt: nowISO() };
                else { const blob = await (await fetch(fileUrl(src))).blob(); const up = await ServerHost.uploadMany([new File([blob], (t.title || 'post').replace(/\\W+/g, '-') + '.jpg', { type: 'image/jpeg' })], { taskId: t.id }); a = { ...up[0], id: uid(), addedAt: nowISO() }; }
                t.attachments.push(a);
                t.reviews = (t.reviews || []).concat({ id: uid(), by: t.assignees[0] || 'u-joe', files: [a.id], at: nowISO(), state: 'approved', decidedAt: nowISO(), decidedBy: currentUserId, round: 1, epoch: t.reviewEpoch || 0 });
                syncCD(t); if (!t.cdStatus) t.cdStatus = 'approved'; t.updatedAt = nowISO(); }
              persist(0); return ts.length; })"""
            ctx, pg, errs = await login('carine@e11-digital.com')
            E = pg.evaluate
            n_ok = await E(APPROVE + "(" + json.dumps(vid) + ")"); await pg.wait_for_timeout(1500)
            await E("Plans.open('pl-smp-p-korea')"); await pg.wait_for_timeout(500)
            total = await E("Plans.planById('pl-smp-p-korea').posts.length")
            check(f'{total} of {total}' in await pg.locator('.pl-banner').inner_text(), f'all {total} posts finished after the CD approves ({n_ok})')
            await ctx.close()
            # Aishwarya exports Round 1
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            await E("Plans.open('pl-smp-p-korea')"); await pg.wait_for_timeout(500)
            await pg.locator('[data-pl="export"]').click(); await pg.wait_for_timeout(400)
            if await pg.locator('.ex-logo[data-logo]').count():
                await pg.locator('.ex-logo[data-logo]').first.click(); await pg.wait_for_timeout(200)
            else:
                async with pg.expect_file_chooser() as fc:
                    await pg.locator('[data-x="upload-logo"]').click()
                await (await fc.value).set_files(os.path.join(ROOT, 'server/src/sample-files/clickup/cu09.jpg')); await pg.wait_for_timeout(1500)
            check(await pg.locator('.ex-cur').count() == 1, 'the logo is picked (from Client Assets › Logos)')
            check(await pg.locator('.ex-row').count() == total, f'{total} finished posts go in')
            await pg.screenshot(path=os.path.join(OUT, 'plans-export.png'))
            await pg.locator('.ex-dlg [data-r="1"]').click()
            await pg.wait_for_selector('.ex-ok', timeout=240000)
            check(await pg.locator('.ex-ok .pl-dl').count() == 2, 'PowerPoint and PDF ready to download')
            await pg.locator('.ex-dlg [data-r="0"]').last.click(); await pg.wait_for_timeout(500)
            check('With the client · Round 1' in await pg.locator('.pl-who').inner_text(), 'With the client · Round 1')
            ex = await E("Plans.planById('pl-smp-p-korea').exports[0]")
            check(await E("!!proj('p-korea').planLogo"), 'the logo is remembered for next time')
            raw = await asyncio.to_thread(lambda: urllib.request.urlopen(urllib.request.Request(BASE + '/__file/' + urllib.parse.quote(ex['pptx']['path']), headers={'X-Dev-Email': EXEC})).read())
            open(os.path.join(OUT, 'round1.pptx'), 'wb').write(raw)
            from pptx import Presentation
            prs = Presentation(os.path.join(OUT, 'round1.pptx'))
            slides = list(prs.slides)
            check(len(slides) == total + 4, f'the deck: cover, month, {total} posts, how to review, thank you ({len(slides)})')
            boxes = [sh for sl in slides for sh in sl.shapes if sh.name.startswith('E11 client comments')]
            check(len(boxes) == total, 'a Client comments box on each post')
            vids = [sh for sl in slides for sh in sl.shapes if sh.name == 'kpop-night.mp4']
            check(len(vids) == 1, 'the reel’s video is in it')
            # the client comments on one post and sends it back
            first = boxes[0]; first.text_frame.text = 'Make the title bigger please'
            prs.save(os.path.join(OUT, 'back.pptx'))
            changed_post = first.name.replace('E11 client comments ', '')
            await pg.locator('[data-pl="reply"]').click(); await pg.wait_for_timeout(300)
            async with pg.expect_file_chooser() as fc:
                await pg.locator('.rp-dlg [data-x="upload"]').click()
            await (await fc.value).set_files(os.path.join(OUT, 'back.pptx')); await pg.wait_for_timeout(2500)
            check(await pg.locator('.rp-row.changes').count() == 1 and await pg.locator('.rp-row.approved').count() == total - 1, 'reading the client’s PowerPoint: 1 change, the rest approved')
            check('Make the title bigger please' in await pg.locator('.rp-note').input_value(), 'their comment is on the post')
            await pg.screenshot(path=os.path.join(OUT, 'plans-reply.png'))
            await pg.locator('.rp-dlg [data-r="1"]').click(); await pg.wait_for_timeout(1500)
            st = await E("db.tasks.filter(t => t.planId === 'pl-smp-p-korea').map(t => [t.planPostId, t.done, statusList(t).find(s => s.id === t.sectionId).name, (t.comments || []).some(c => /Client’s changes/.test(c.text))])")
            done_n = len([x for x in st if x[1]]); back = [x for x in st if x[0] == changed_post][0]
            check(done_n == total - 1, f'{total - 1} approved posts completed ({done_n})')
            check(back[1] is False and back[2] == 'In progress' and back[3], f'the changed post’s task is back In progress with the client’s comment ({back})')
            check(await pg.locator(f'.pl-chip[data-id="{changed_post}"]').count() == 1 and 'Changes' in await pg.locator(f'.pl-chip[data-id="{changed_post}"]').inner_text(), 'the calendar shows the client’s changes')
            check(await pg.locator('[data-pl="client-all"]').count() == 0, 'no Client approved for the whole plan while a post is being changed')
            await pg.wait_for_timeout(800)
            cl = await asyncio.to_thread(lambda: next(c for c in api(ADMIN, '/api/snapshot')['data']['clients'] if c['id'] == next(p['clientId'] for p in api(ADMIN, '/api/snapshot')['data']['projects'] if p['id'] == 'p-korea')))
            pf = next((f for f in cl['folders'] if f.get('planOf') == 'pl-smp-p-korea'), None)
            check(pf is not None and pf['name'].startswith('Content plan – '), 'Client Assets › Files › Content plan – <month>')
            inside = [f for f in cl['folders'] if pf and f.get('parentId') == pf['id'] and f.get('finalOf')]
            check(len(inside) == total - 1, f'a final-files folder per approved task inside it ({len(inside)})')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # the CD approves the new version; Round 2; the client approves
            ctx, pg, errs = await login('carine@e11-digital.com')
            E = pg.evaluate
            await E(APPROVE + "(null)"); await pg.wait_for_timeout(1500)
            await ctx.close()
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            await E("Plans.open('pl-smp-p-korea')"); await pg.wait_for_timeout(500)
            check('Finished' in await pg.locator(f'.pl-chip[data-id="{changed_post}"]').inner_text(), 'the new version is in the plan once the CD approves it')
            await pg.locator('[data-pl="export"]').click(); await pg.wait_for_timeout(300)
            check(await pg.locator('.ex-cur').count() == 1, 'Round 2 uses the same logo')
            await pg.locator('.ex-dlg [data-r="1"]').click()
            await pg.wait_for_selector('.ex-ok', timeout=240000)
            await pg.locator('.ex-dlg [data-r="0"]').last.click(); await pg.wait_for_timeout(500)
            check('Round 2' in await pg.locator('.pl-who').inner_text(), 'With the client · Round 2')
            check(await pg.locator('[data-pl="client-all"]').count() == 1, 'every post finished: Client approved')
            await pg.locator('[data-pl="client-all"]').click(); await pg.wait_for_timeout(300)
            await pg.locator('.dialog [data-r="1"]').click(); await pg.wait_for_timeout(2500)
            check(await E("!!Plans.planById('pl-smp-p-korea').clientApproved") and await E("db.tasks.filter(t => t.planId === 'pl-smp-p-korea').every(t => t.done)"), 'the content plan is client approved and every task complete')
            await pg.screenshot(path=os.path.join(OUT, 'plans-client-approved.png'))
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            await asyncio.sleep(1)
            snap = await asyncio.to_thread(lambda: api(ADMIN, '/api/snapshot')['data'])
            pl = next(p for p in snap['plans'] if p['id'] == 'pl-smp-p-korea')
            cl = next(c for c in snap['clients'] if c['id'] == next(p['clientId'] for p in snap['projects'] if p['id'] == 'p-korea'))
            pf = next(f for f in cl['folders'] if f.get('planOf') == 'pl-smp-p-korea')
            check(pl.get('filedDeck') and any(f['name'].endswith('Round 2.pptx') for f in pf['files']), 'the approved PowerPoint (Round 2) is filed in the plan’s folder')
            check(len([f for f in cl['folders'] if f.get('parentId') == pf['id'] and f.get('finalOf')]) == total, 'with a final-files folder for every task')
            # notifications and their links: the AM (tasks made), Aishwarya (approved), Joe (the client's changes)
            ctx, pg, errs = await login(AM)
            E = pg.evaluate
            await E("go('inbox')"); await pg.wait_for_timeout(400)
            check(await E("myNotifs().filter(n => n.kind === 'plan-tasks').length") >= 1, 'the AM was told about the tasks made from the skeleton')
            row = pg.locator('.ib-row', has_text='KOREA 360 · November 2026')
            check(await row.count() == 1, 'one Inbox row for the plan’s updates')
            await row.first.click(); await pg.wait_for_timeout(500)
            check(await E("ui.view") == 'plan' and await E("ui.planId") == 'pl-smp-p-korea', 'it opens the plan')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            ctx, pg, errs = await login(EXEC)
            E = pg.evaluate
            await E("go('inbox'); ui.inboxTab = 'all'; renderContent()"); await pg.wait_for_timeout(400)
            ok_n = await E("myNotifs().filter(n => n.kind === 'plan-ok').length")
            check(ok_n == 1, 'Aishwarya has the approval update')
            await E("openNotif(myNotifs().find(n => n.kind === 'plan-ok').id)"); await pg.wait_for_timeout(400)
            check(await E("ui.view") == 'plan', 'the approval update opens the plan')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            joe = await asyncio.to_thread(lambda: [n for n in api('joe@e11-digital.com', '/api/snapshot')['data']['notifications'] if n['kind'] == 'comment' and 'client’s changes' in n['text']])
            am_c = await asyncio.to_thread(lambda: [n for n in api(AM, '/api/snapshot')['data']['notifications'] if n['kind'] == 'comment' and 'client’s changes' in n['text']])
            check(len(am_c) >= 1, f'the task’s people (here the AM) hear about the client’s changes ({len(am_c)})')
            # phone: the Inbox row and a tapped phone notification open the plan
            ctx, pg, errs = await login(MGR, (390, 844))
            E = pg.evaluate
            await pg.click('[data-ph-tab="inbox"]'); await pg.wait_for_timeout(500)
            prow = pg.locator('.ph-ib', has_text='Content')
            check(await prow.count() >= 1, 'phone Inbox shows the content plan updates')
            await prow.first.click(); await pg.wait_for_timeout(600)
            check(await E("ui.view") == 'plan' and await E("ui.planId") == 'pl-smp-p-korea', 'phone: tapping it opens the plan')
            nid = await E("myNotifs().find(n => n.planId)?.id")
            await pg.goto(f'/?n={nid}&k=plan&pl=pl-smp-p-korea'); await pg.wait_for_timeout(4500)
            check(await E("ui.view") == 'plan' and await E("ui.planId") == 'pl-smp-p-korea', 'a tapped phone notification opens the plan')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # the Social Manager is told; the update opens the plan
            ctx, pg, errs = await login(MGR)
            E = pg.evaluate
            await E("go('inbox'); ui.inboxTab = 'all'; renderContent()"); await pg.wait_for_timeout(400)
            row = pg.locator('.ib-row', has_text='Content plan')
            check(await row.count() >= 1 and 'client’s approval' in await row.first.inner_text(), 'Amelia: the client approved the content plan')
            await row.first.click(); await pg.wait_for_timeout(500)
            check(await E("ui.view") == 'plan' and await pg.locator('.pl-banner.approved').count() == 1, 'the update opens the approved plan')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # ---------- the Account Manager: view only ----------
            ctx, pg, errs = await login(AM)
            E = pg.evaluate
            await E("go('plans')"); await pg.wait_for_timeout(400)
            check(await pg.locator('[data-pl="start"]').count() == 0, 'an AM can’t start plans')
            await E("ui.planId='pl-smp-p-cpt'; go('plan')"); await pg.wait_for_timeout(400)
            check(await pg.locator('.pl-ro').count() == 1 and await pg.locator('[data-pl="new-post"]').count() == 0, 'an AM sees a plan view only')
            await pg.locator('.pl-chip').first.click(); await pg.wait_for_timeout(300)
            check(await pg.locator('#pp-title').get_attribute('readonly') is not None and await pg.locator('.pl-dlg [type="submit"]').count() == 0, 'a post opens read-only')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # ---------- the Social Manager: every client ----------
            ctx, pg, errs = await login(MGR)
            E = pg.evaluate
            await E("go('plans')"); await pg.wait_for_timeout(400)
            n = await pg.locator('.pl-row:not(.tt-head)').count()
            check(n == await E("db.projects.filter(p => !p.inactive).length"), f'the Social Manager sees every client ({n})')
            check(await pg.locator('[data-pl="start"]').count() >= 5, 'and can start any of them')
            check(await E("designationOf(me())") == 'Social Manager', 'shown as Social Manager')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # ---------- Admin: the Social Manager pick and a client's Social people ----------
            ctx, pg, errs = await login(ADMIN)
            E = pg.evaluate
            await E("go('team')"); await pg.wait_for_timeout(400)
            check('Social Manager' in await pg.locator('.tt-row', has_text='Amelia').inner_text(), 'the Team page shows the Social Manager')
            await pg.locator('.tt-row', has_text='Aishwarya').locator('[data-action="member-menu"]').click(); await pg.wait_for_timeout(200)
            check(await pg.locator('.menu .mi', has_text='Social Manager').count() == 1, 'Social people have a Social Manager option')
            await pg.mouse.click(700, 20); await pg.wait_for_timeout(200)
            await E("ClientList.dialog(proj('p-cpt'))"); await pg.wait_for_timeout(300)
            check('Maria' in await pg.locator('#cd-socials').inner_text(), 'the client pop-up lists its Social people')
            await pg.locator('[data-cd="add-social"]').click(); await pg.wait_for_timeout(200)
            check(await pg.locator('[data-cd-am]').count() >= 2, 'picking from the Social people')
            await pg.locator('[data-cd-am]', has_text='Aishwarya').click(); await pg.wait_for_timeout(200)
            await pg.keyboard.press('Escape'); await pg.wait_for_timeout(150)
            await pg.screenshot(path=os.path.join(OUT, 'plans-client-popup.png'))
            await pg.locator('.cd-dlg [type="submit"]').click(); await pg.wait_for_timeout(700)
            socials = await asyncio.to_thread(lambda: next(p for p in api(ADMIN, '/api/snapshot')['data']['projects'] if p['id'] == 'p-cpt').get('socials'))
            check(socials == ['u-new', 'u-aish'], f'saved: {socials}')
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            # ---------- phone ----------
            ctx, pg, errs = await login(EXEC, (390, 844))
            await pg.click('[data-ph-tab="me"]'); await pg.wait_for_timeout(400)
            await pg.locator('.ph-row[data-nav="plans"]').click(); await pg.wait_for_timeout(500)
            check(await pg.locator('.pl-row:not(.tt-head)').count() == 3, 'phone: the list (CPT was just added for her)')
            await pg.screenshot(path=os.path.join(OUT, 'plans-phone-list.png'))
            await pg.locator('.pl-row', has_text='KOREA 360').click(); await pg.wait_for_timeout(500)
            check(await pg.locator('.pl-post').count() >= 6, 'phone: the plan as a list by day')
            await pg.screenshot(path=os.path.join(OUT, 'plans-phone.png'))
            await pg.locator('.pl-post').first.click(); await pg.wait_for_timeout(400)
            bb = await pg.locator('.pl-dlg').bounding_box()
            check(bb and bb['width'] >= 389 and bb['height'] <= 845, f'phone: a post opens full screen {bb}')
            await pg.screenshot(path=os.path.join(OUT, 'plans-phone-post.png'))
            check(not errs, f'no page errors {errs}')
            await ctx.close()
            await b.close()
    finally:
        srv.terminate()
    print('FAILS', len(fails), fails)
asyncio.run(main())
